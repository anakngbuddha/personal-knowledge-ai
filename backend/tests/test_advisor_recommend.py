"""4.2 Gap analysis and recommendation.

Offline and deterministic. The engine is a pure function over rows, so every claim in
this file is checked without a database, a model, or an API key. The export tests open
the produced files back up rather than only checking that bytes came out.
"""

from __future__ import annotations

import io
import uuid
import zipfile
from types import SimpleNamespace

from app.advisor.brief import CustomerBrief, RoomRequirement
from app.advisor.export import to_docx, to_pptx
from app.advisor.recommend import (
    Citation,
    recommend_from_rows,
    tokens,
)
from app.db.models import RelationType


def product(
    name: str,
    *,
    vendor: str = "Jabra",
    category: str = "headsets",
    description: str = "",
    deployment: str = "cloud",
    lifecycle: str = "GA",
    curation: str = "confirmed",
    demo: bool = False,
):
    return SimpleNamespace(
        id=uuid.uuid4(),
        name=name,
        slug=name.lower().replace(" ", "-"),
        vendor=vendor,
        category=category,
        description=description,
        tier="Core",
        deployment_model=deployment,
        lifecycle_status=lifecycle,
        curation_status=curation,
        is_demo=demo,
        collateral_document_ids=[],
    )


def edge(source, target, relation: str, evidence: str = "the datasheet says so"):
    return SimpleNamespace(
        source_product_id=str(source.id),
        target_product_id=str(target.id),
        relation_type=relation,
        evidence=evidence,
        confidence=0.9,
    )


BRIEF = CustomerBrief(
    customer="Northwind Logistics",
    industry="logistics",
    rooms=[
        RoomRequirement(room_type="huddle room", count=12, platform="Microsoft Teams"),
        RoomRequirement(room_type="contact center", seats=40),
        RoomRequirement(room_type="boardroom", count=1),
    ],
    cloud_needs=["backup", "disaster recovery"],
    platforms=["Microsoft Teams"],
    budget="$180,000",
)

CATALOGUE = [
    product(
        "Poly Studio X50",
        vendor="Poly",
        category="conferencing",
        description="All-in-one video bar for a huddle room, certified for Microsoft Teams.",
    ),
    product(
        "Jabra Engage 50",
        vendor="Jabra",
        category="headsets",
        description="Professional contact center headset for agents on Microsoft Teams.",
    ),
    product(
        "Shure MXA920",
        vendor="Shure",
        category="audio",
        description="Ceiling array microphone for a boardroom.",
    ),
    product(
        "Huawei Cloud Backup",
        vendor="Huawei",
        category="cloud",
        description="Cloud backup and disaster recovery for enterprise workloads.",
    ),
    product(
        "Legacy Bridge 100",
        vendor="Poly",
        category="conferencing",
        description="Huddle room bridge for Microsoft Teams.",
        lifecycle="EOL",
    ),
]


def build(brief=BRIEF, products=None, **kwargs):
    return recommend_from_rows(
        brief=brief,
        products=CATALOGUE if products is None else products,
        **kwargs,
    )


# -- scoring and coverage --------------------------------------------------


def test_tokens_drop_noise_words():
    assert "the" not in tokens("The boardroom needs a ceiling microphone")
    assert {"boardroom", "ceiling", "microphone"} <= tokens(
        "The boardroom needs a ceiling microphone"
    )


def test_each_requirement_gets_products_from_the_product_list():
    result = build()
    by_requirement = {line.requirement: line for line in result.lines}

    huddle = by_requirement["12 huddle rooms on Microsoft Teams"]
    assert "Poly Studio X50" in [pick.name for pick in huddle.picks]

    agents = by_requirement["40-seat contact center"]
    assert "Jabra Engage 50" in [pick.name for pick in agents.picks]

    board = by_requirement["1 boardroom"]
    assert "Shure MXA920" in [pick.name for pick in board.picks]


def test_nothing_outside_the_product_list_is_ever_recommended():
    names = {pick.name for pick in build().picks}
    assert names <= {item.name for item in CATALOGUE}


def test_end_of_life_and_sample_products_are_not_quoted():
    result = build(
        products=[
            *CATALOGUE,
            product("Sample Bar", category="conferencing", description="huddle room", demo=True),
            product(
                "Draft Bar",
                category="conferencing",
                description="huddle room",
                curation="suggested",
            ),
        ]
    )
    names = {pick.name for pick in result.picks}
    assert "Legacy Bridge 100" not in names
    assert "Sample Bar" not in names
    assert "Draft Bar" not in names


def test_a_room_and_its_platform_both_raise_the_score():
    result = build()
    huddle = next(line for line in result.lines if line.requirement.startswith("12 huddle"))
    pick = next(item for item in huddle.picks if item.name == "Poly Studio X50")
    assert any("huddle room" in reason for reason in pick.reasons)
    assert any("Microsoft Teams" in reason for reason in pick.reasons)


def test_a_platform_name_alone_is_not_a_match():
    """Half a catalogue is "certified for Teams". That cannot put a headset in a room line."""
    huddle = next(
        line for line in build().lines if line.requirement.startswith("12 huddle")
    )
    assert [pick.name for pick in huddle.picks] == ["Poly Studio X50"]


def test_the_engine_is_deterministic():
    assert build().as_dict() == build().as_dict()


# -- gaps -----------------------------------------------------------------


def test_a_requirement_nothing_covers_is_reported_as_a_gap_not_padded():
    brief = CustomerBrief(
        customer="Acme",
        cloud_needs=["submarine cable landing station"],
    )
    result = build(brief=brief)
    assert not result.lines
    assert any("submarine" in gap for gap in result.gaps)


def test_an_empty_product_list_says_so_instead_of_inventing_a_bundle():
    result = build(products=[])
    assert not result.lines
    assert "product list" in result.as_markdown()


# -- provenance -----------------------------------------------------------


def test_a_product_with_a_document_behind_it_cites_it():
    poly = CATALOGUE[0]
    result = build(
        evidence={
            str(poly.id): [
                Citation(
                    label="studio-x50.pdf#p2",
                    document_title="Poly Studio X50 datasheet",
                    quote="Certified for Microsoft Teams Rooms.",
                    page_number=2,
                )
            ]
        }
    )
    pick = next(item for item in result.picks if item.name == "Poly Studio X50")
    assert pick.grounded
    assert pick.provenance == "your documents"
    assert "Poly Studio X50 datasheet, page 2" in pick.why()


def test_a_product_with_no_document_is_labelled_general_knowledge_and_flagged():
    result = build()
    pick = next(item for item in result.picks if item.name == "Shure MXA920")
    assert not pick.grounded
    assert "general product knowledge" in pick.why()
    assert any("verify" in item.lower() or "confirm" in item.lower() for item in result.assumptions)


def test_general_knowledge_is_never_presented_as_a_document():
    markdown = build().as_markdown()
    assert "not from a document you uploaded" in markdown


# -- the map --------------------------------------------------------------


def test_a_clash_on_the_map_drops_one_side_and_says_why():
    poly, jabra = CATALOGUE[0], CATALOGUE[1]
    result = build(
        edges=[edge(poly, jabra, RelationType.CONFLICTS_WITH, "cannot share the same DSP")]
    )
    names = {pick.name for pick in result.picks}
    assert len(names & {"Poly Studio X50", "Jabra Engage 50"}) == 1
    assert result.conflicts
    assert "cannot share the same DSP" in result.conflicts[0]["reason"]


def test_a_prerequisite_outside_the_bundle_is_reported():
    """It is in the product list, but nothing the customer said would surface it."""
    poly = CATALOGUE[0]
    injector = product("Orbit PoE Injector", vendor="Orbit", category="networking")
    result = build(
        products=[*CATALOGUE, injector],
        edges=[edge(poly, injector, RelationType.REQUIRES, "needs PoE+ at the wall plate")],
    )
    assert "Orbit PoE Injector" not in {pick.name for pick in result.picks}
    needed = {row["needed"]: row for row in result.also_needed}
    assert "Orbit PoE Injector" in needed
    assert needed["Orbit PoE Injector"]["in_product_list"] is True
    assert needed["Orbit PoE Injector"]["product"] == "Poly Studio X50"


def test_a_prerequisite_missing_from_the_product_list_is_called_out():
    poly = CATALOGUE[0]
    stranger = product("Unlisted Switch", vendor="Other", category="networking")
    result = build(
        products=CATALOGUE,  # the prerequisite is deliberately not in the list
        edges=[edge(poly, stranger, RelationType.REQUIRES, "needs a PoE+ switch")],
    )
    rows = [row for row in result.also_needed if not row["in_product_list"]]
    assert rows
    assert any("not in your product list" in item for item in result.assumptions)


def test_upsell_and_cross_sell_come_from_the_map():
    poly = CATALOGUE[0]
    bigger = product("Poly Studio X70", vendor="Poly", category="conferencing")
    extra = product("Poly TC10 controller", vendor="Poly", category="conferencing")
    result = build(
        products=[*CATALOGUE, bigger, extra],
        edges=[
            edge(poly, bigger, RelationType.UPSELL_TO),
            edge(poly, extra, RelationType.CROSS_SELL),
        ],
    )
    growth = {row["product"]: row["kind"] for row in result.growth}
    assert growth.get("Poly Studio X70") == "upsell"
    assert growth.get("Poly TC10 controller") == "cross-sell"


def test_only_approved_relationships_reach_the_answer():
    """The caller filters to approved rows; passing none must yield no graph claims."""
    result = build(edges=[])
    assert result.conflicts == []
    assert result.also_needed == []
    assert result.growth == []


# -- exclusions, questions, assumptions -----------------------------------


def test_a_vendor_the_customer_ruled_out_is_dropped_and_recorded():
    brief = CustomerBrief(
        customer="Northwind",
        rooms=[RoomRequirement(room_type="huddle room", count=12, platform="Microsoft Teams")],
        exclusions=["Poly"],
    )
    result = build(brief=brief)
    assert "Poly Studio X50" not in {pick.name for pick in result.picks}
    assert any(row["product"] == "Poly Studio X50" for row in result.excluded)


def test_an_excluded_vendor_is_not_offered_as_an_upsell_either():
    poly, jabra = CATALOGUE[0], CATALOGUE[1]
    brief = CustomerBrief(
        customer="Northwind",
        rooms=[RoomRequirement(room_type="contact center", seats=40)],
        exclusions=["Poly"],
    )
    result = build(brief=brief, edges=[edge(jabra, poly, RelationType.UPSELL_TO)])
    assert "Poly Studio X50" not in {row["product"] for row in result.growth}


def test_missing_brief_fields_become_questions_for_the_customer():
    brief = CustomerBrief(
        customer="Acme",
        rooms=[RoomRequirement(room_type="boardroom")],
    )
    questions = build(brief=brief).questions
    assert any("budget" in item.lower() for item in questions)
    assert any("live" in item.lower() for item in questions)
    assert any("boardroom" in item.lower() for item in questions)


def test_pricing_is_never_included_and_is_always_flagged():
    result = build()
    assert any("price" in item.lower() for item in result.assumptions)


# -- the written answer ----------------------------------------------------


def test_the_answer_restates_the_requirements_first():
    markdown = build().as_markdown()
    assert markdown.startswith("## What Northwind Logistics asked for")
    assert "12 huddle rooms on Microsoft Teams" in markdown


def test_the_answer_is_prose_plus_one_compact_table():
    markdown = build().as_markdown()
    assert "### At a glance" in markdown
    assert markdown.count("| Requirement | Product |") == 1


def test_the_answer_uses_plain_words_only():
    markdown = build().as_markdown()
    for banned in ("chunk", "ingest", " MCP", " RLS", "dossier", "collateral", "hybrid search"):
        assert banned.lower() not in markdown.lower()


def test_the_payload_carries_every_section_the_screen_needs():
    payload = build().as_dict()
    for key in (
        "requirements_restated",
        "recommendations",
        "conflicts",
        "also_needed",
        "gaps",
        "growth",
        "questions_to_ask",
        "assumptions_to_verify",
        "table",
    ):
        assert key in payload


# -- exports ---------------------------------------------------------------


def _docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def test_word_export_opens_and_carries_the_recommendation():
    from docx import Document

    result = build()
    data = to_docx(result)
    document = Document(io.BytesIO(data))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Northwind Logistics" in document.paragraphs[0].text
    assert "Poly Studio X50" in text
    assert document.tables


def test_word_export_keeps_the_general_knowledge_caveat():
    assert "general product knowledge" in _docx_text(to_docx(build())).lower()


def test_powerpoint_export_opens_with_a_slide_per_requirement():
    from pptx import Presentation

    result = build()
    deck = Presentation(io.BytesIO(to_pptx(result)))
    titles = [
        slide.shapes.title.text
        for slide in deck.slides
        if slide.shapes.title is not None
    ]
    assert titles[0].startswith("Recommendation for Northwind Logistics")
    assert "What they asked for" in titles
    assert any("huddle" in title for title in titles)
    assert "Verify before you send this" in titles


def test_an_empty_recommendation_still_exports_without_crashing():
    empty = build(products=[])
    assert len(to_docx(empty)) > 0
    assert len(to_pptx(empty)) > 0
