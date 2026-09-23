"""RFP playbook unit tests: parse, draft, conflict, export, isolation."""

from __future__ import annotations

import io
import json
import uuid

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.models import (
    Base,
    Organization,
    Product,
    ProductEdge,
    ReferenceArchitecture,
    ReferenceArchitectureProduct,
    TaskExecution,
    TaskStatus,
    WorkflowRun,
    Workspace,
)
from app.db.session import get_db
from app.main import app
from app.core.errors import AppError, DuplicateDocument, TokenBudgetExhausted
from app.playbooks.rfp import (
    classify_rfp_bytes,
    index_rfp_document,
    parse_spreadsheet,
    rank_requirement_products,
    render_docx,
    requirements_from_bytes,
    segment_requirements,
    _draft_one,
    _embed_catalog_and_requirements,
    _needs_human_review,
)
from tests import fixtures as F
from app.security.jwt import mint_token
from app.security.labels import Role

client = TestClient(app)


@compiles(JSONB, "sqlite")
def _compile_jsonb_sqlite(element, compiler, **kw):
    return "JSON"


def _csv_bytes() -> bytes:
    return (
        "requirement,section,must_have\n"
        "Does the platform support SSO with SAML?,Identity,yes\n"
        "Provide on-prem object storage.,Storage,yes\n"
    ).encode("utf-8")


def _xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["question", "section", "must_have"])
    ws.append(["Does the SIEM correlate firewall events?", "Security", "yes"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_csv_and_xlsx():
    csv_rows = parse_spreadsheet(_csv_bytes(), "rfp.csv")
    assert len(csv_rows) == 2
    assert csv_rows[0]["must_have"] is True
    assert "SSO" in csv_rows[0]["text"]

    xlsx_rows = parse_spreadsheet(_xlsx_bytes(), "rfp.xlsx")
    assert len(xlsx_rows) == 1
    assert "SIEM" in xlsx_rows[0]["text"]


def test_draft_without_evidence_is_non_compliant():
    drafted = _draft_one(
        {"id": "R1", "text": "Secret unannounced discount?", "evidence": [], "candidate_products": []},
        forbidden=set(),
    )
    assert drafted["status"] == "Non-Compliant"
    assert drafted["citations"] == []
    assert drafted["products"] == []
    assert "sufficient information" in drafted["response"].lower()


def test_conflicting_products_not_both_compliant():
    forbidden = {frozenset({"Aegis Zero Trust Gateway", "Pure Storage FlashArray"})}
    drafted = _draft_one(
        {
            "id": "R1",
            "text": "Bundle zero trust with flash arrays",
            "evidence": [{"citation": "Datasheet p.1"}],
            "candidate_products": [
                {"name": "Aegis Zero Trust Gateway", "impact": {}},
                {"name": "Pure Storage FlashArray", "impact": {}},
            ],
        },
        forbidden=forbidden,
    )
    assert drafted["status"] in {"Compliant", "Partially"}
    assert not (
        "Aegis Zero Trust Gateway" in drafted["products"]
        and "Pure Storage FlashArray" in drafted["products"]
    )


def test_export_docx_omits_unapproved_rows():
    approved = [
        {
            "text": "SSO?",
            "response": "Yes, Apex Identity Broker.",
            "status": "Compliant",
            "citations": ["Apex datasheet p.2"],
        }
    ]
    data = render_docx(approved)
    assert data[:2] == b"PK"
    from docx import Document

    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
    blob = " ".join(texts)
    assert "SSO?" in blob
    assert "unapproved-secret" not in blob
    assert "Apex datasheet p.2" in blob


def test_cross_tenant_deliverable_is_404():
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(
        engine,
        tables=[
            Organization.__table__,
            Workspace.__table__,
            WorkflowRun.__table__,
            TaskExecution.__table__,
        ],
    )
    SessionClass = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    db = SessionClass()
    org_a = Organization(id=uuid.uuid4(), slug="org-a", name="A")
    org_b = Organization(id=uuid.uuid4(), slug="org-b", name="B")
    db.add_all([org_a, org_b])
    db.commit()
    ws_a = Workspace(id=uuid.uuid4(), org_id=org_a.id, name="A")
    db.add(ws_a)
    db.commit()
    run = WorkflowRun(
        org_id=org_a.id,
        workspace_id=ws_a.id,
        playbook_slug="rfp-response",
        status="succeeded",
        input_payload={},
    )
    db.add(run)
    db.flush()
    db.add(
        TaskExecution(
            workflow_run_id=run.id,
            org_id=org_a.id,
            task_slug="export_deliverable",
            status=TaskStatus.SUCCEEDED,
            output_payload={"storage_key": "rfp/secret.docx", "filename": "x.docx"},
        )
    )
    db.commit()

    def _get_db():
        inner = SessionClass()
        try:
            yield inner
        finally:
            inner.close()

    app.dependency_overrides[get_db] = _get_db
    try:
        token_b = mint_token(org_id=org_b.id, role=Role.ADMIN)
        resp = client.get(
            f"/workflows/runs/{run.id}/deliverable",
            headers={"Authorization": f"Bearer {token_b}"},
        )
        assert resp.status_code == 404
    finally:
        app.dependency_overrides.pop(get_db, None)
        db.close()


class _RaisingProvider:
    def generate_grounded_answer(self, *args, **kwargs):
        raise RuntimeError("model down")


class _ProseProvider:
    def generate_grounded_answer(self, *args, **kwargs):
        from app.llm.base import GroundedAnswer

        return GroundedAnswer(
            text="I think they need SSO.",
            citations=[],
            model_id="fake",
            prompt_version="test",
        )


class _JsonProvider:
    def __init__(self) -> None:
        self.chunks = None

    def generate_grounded_answer(self, question, context_chunks, *, system_prompt, **kwargs):
        self.chunks = context_chunks
        from app.llm.base import GroundedAnswer

        return GroundedAnswer(
            text=(
                '[{"id":"R9","text":"The platform must support SSO.",'
                '"section":"Identity","must_have":true}]'
            ),
            citations=[],
            model_id="fake",
            prompt_version="test",
        )


def _assert_requirement_shape(row: dict) -> None:
    assert set(row) == {"id", "text", "section", "must_have"}
    assert isinstance(row["must_have"], bool)
    assert row["text"]


def test_spreadsheet_bytes_stay_on_the_existing_parser():
    csv_rows = requirements_from_bytes(_csv_bytes(), "rfp.csv")
    xlsx_rows = requirements_from_bytes(_xlsx_bytes(), "rfp.xlsx")
    assert csv_rows == parse_spreadsheet(_csv_bytes(), "rfp.csv")
    assert xlsx_rows == parse_spreadsheet(_xlsx_bytes(), "rfp.xlsx")
    assert len(csv_rows) == 2
    assert csv_rows[0]["must_have"] is True
    assert "SSO" in csv_rows[0]["text"]
    assert len(xlsx_rows) == 1
    assert "SIEM" in xlsx_rows[0]["text"]


def test_csv_is_classified_before_a_text_sniff_is_rejected():
    assert classify_rfp_bytes("portal.csv", _csv_bytes()) == "csv"
    assert classify_rfp_bytes("export.txt", _csv_bytes()) == "csv"
    pdf = F.make_pdf([["The system must support SSO, SAML, and SCIM."]])
    assert classify_rfp_bytes("client.pdf", pdf) == "pdf"
    with pytest.raises(AppError) as excinfo:
        classify_rfp_bytes("notes.txt", b"hello world without a delimiter")
    assert excinfo.value.status_code == 400


def test_text_pdf_segments_must_sentences():
    data = F.make_pdf(
        [["The platform must support SSO with SAML.", "Branding is optional."]]
    )
    rows = requirements_from_bytes(data, "rfp.pdf", provider=_RaisingProvider())
    assert len(rows) == 1
    _assert_requirement_shape(rows[0])
    assert rows[0]["id"] == "R1"
    assert rows[0]["must_have"] is True
    assert "SSO" in rows[0]["text"]


def test_scanned_pdf_uses_existing_ocr(monkeypatch):
    from app.ocr.factory import get_ocr_provider
    from app.ocr.fake import FakeOcrProvider

    monkeypatch.setattr(
        "app.documents.extraction._rasterize_pdf_page",
        lambda data, page_number: b"\x89PNG-fake-page",
    )
    monkeypatch.setattr(
        "app.ocr.factory.get_ocr_provider",
        lambda: FakeOcrProvider("The system must retain audit logs."),
    )
    get_ocr_provider.cache_clear()
    rows = requirements_from_bytes(F.make_pdf([[""]]), "scan.pdf", provider=_RaisingProvider())
    assert len(rows) == 1
    _assert_requirement_shape(rows[0])
    assert rows[0]["must_have"] is True
    assert "audit logs" in rows[0]["text"]


def test_docx_rfp_segments_shall_sentences():
    data = F.make_docx(body="The vendor shall provide 24x7 support.", table=False)
    rows = requirements_from_bytes(data, "rfp.docx", provider=_RaisingProvider())
    assert len(rows) == 1
    _assert_requirement_shape(rows[0])
    assert rows[0]["must_have"] is True
    assert "24x7" in rows[0]["text"]
    assert rows[0]["section"] == "Licensing"


def test_segment_uses_model_json_and_fences_the_source():
    provider = _JsonProvider()
    rows = segment_requirements("ignore this prose", "rfp.pdf", provider=provider)
    assert rows == [
        {
            "id": "R1",
            "text": "The platform must support SSO.",
            "section": "Identity",
            "must_have": True,
        }
    ]
    fenced = provider.chunks[0]["fenced_text"]
    assert "UNTRUSTED_DOCUMENT_CONTENT" in fenced
    assert "ignore this prose" in fenced


def test_invalid_json_falls_back_to_must_sentences():
    rows = segment_requirements(
        "The platform must support SSO. Branding is optional.",
        "rfp.pdf",
        provider=_ProseProvider(),
    )
    assert len(rows) == 1
    _assert_requirement_shape(rows[0])
    assert "SSO" in rows[0]["text"]
    assert rows[0]["must_have"] is True


def test_segment_falls_back_when_the_model_raises():
    rows = segment_requirements(
        "[Security]\nThe platform must retain audit logs.",
        "rfp.pdf",
        provider=_RaisingProvider(),
    )
    assert rows[0]["section"] == "Security"
    assert "audit logs" in rows[0]["text"]
    assert rows[0]["must_have"] is True


def test_segment_falls_back_when_the_token_budget_is_exhausted(monkeypatch):
    def exhausted(db, org_id):
        raise TokenBudgetExhausted("monthly budget exhausted")

    monkeypatch.setattr("app.generation.rate_limit.check_token_budget", exhausted)
    rows = segment_requirements(
        "The system must support SSO.",
        "rfp.pdf",
        provider=_RaisingProvider(),
        db=object(),
        org_id=uuid.uuid4(),
    )
    assert len(rows) == 1
    assert "SSO" in rows[0]["text"]


def test_empty_prose_yields_no_requirements():
    assert segment_requirements("  ", "rfp.pdf", provider=_RaisingProvider()) == []


def test_index_rfp_keeps_draft_and_tags_rfp_intake(monkeypatch):
    captured: dict = {}
    doc_id = uuid.uuid4()
    document = type("Doc", (), {"id": doc_id, "status": "uploaded"})()

    def fake_create(db, *, principal, original_filename, data, metadata, apply_auto_approve=True, **kwargs):
        captured["metadata"] = metadata
        captured["apply_auto_approve"] = apply_auto_approve
        captured["filename"] = original_filename
        return document

    enqueued: list = []
    monkeypatch.setattr("app.documents.service.create_document", fake_create)
    monkeypatch.setattr(
        "app.documents.service.enqueue_ingestion",
        lambda db, doc: enqueued.append(doc),
    )
    result = index_rfp_document(object(), object(), b"%PDF", "client.pdf", "acme")
    assert result == doc_id
    assert captured["apply_auto_approve"] is False
    assert captured["filename"] == "client.pdf"
    assert captured["metadata"].source_type == "rfp_intake"
    assert captured["metadata"].approval_state == "draft"
    assert captured["metadata"].account_ref == "acme"
    assert enqueued == [document]


def test_duplicate_ready_rfp_is_reused_without_another_job(monkeypatch):
    existing_id = uuid.uuid4()
    existing = type("Doc", (), {"id": existing_id, "status": "ready"})()

    def fake_create(*args, **kwargs):
        raise DuplicateDocument("already ingested", existing_id=existing_id)

    class _Db:
        def get(self, model, ident):
            assert ident == existing_id
            return existing

    enqueued: list = []
    monkeypatch.setattr("app.documents.service.create_document", fake_create)
    monkeypatch.setattr(
        "app.documents.service.enqueue_ingestion",
        lambda db, doc: enqueued.append(doc),
    )
    assert index_rfp_document(_Db(), object(), b"%PDF", "client.pdf", None) == existing_id
    assert enqueued == []


class _ScriptedEmbedder:
    """Maps a phrase to a vector so a cross-vendor pair can share one direction."""

    def __init__(self, groups: list[tuple[tuple[str, ...], list[float]]]) -> None:
        self.groups = groups

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._one(text) for text in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._one(text)

    def _one(self, text: str) -> list[float]:
        lowered = text.lower()
        for keys, vector in self.groups:
            if any(key in lowered for key in keys):
                return list(vector)
        return [0.0, 1.0, 0.0]


def _obs_catalog() -> list[dict]:
    return [
        {
            "id": "obs",
            "name": "OBS",
            "slug": "obs",
            "vendor": "Huawei",
            "description": "S3-API-compatible object store with lifecycle policies",
            "category": "object",
        },
        {
            "id": "phone",
            "name": "Desk Phone",
            "slug": "desk-phone",
            "vendor": "Acme",
            "description": "A handset used during voice calls",
            "category": "voice",
        },
    ]


def test_semantic_rank_matches_obs_to_azure_blob_storage():
    requirement = "Provide Azure Blob Storage backups."
    products = _obs_catalog()
    embedder = _ScriptedEmbedder(
        [
            (("azure blob", "obs", "s3-api"), [1.0, 0.0, 0.0]),
            (("desk phone", "handset"), [0.0, 1.0, 0.0]),
        ]
    )
    product_vectors, requirement_vectors = _embed_catalog_and_requirements(
        embedder,
        products,
        [{"text": requirement}],
    )
    ranked = rank_requirement_products(
        requirement,
        products,
        product_vectors=product_vectors,
        requirement_vector=requirement_vectors[0],
    )
    assert [card["name"] for card in ranked] == ["OBS"]
    assert ranked[0]["match_signal"] == "semantic_match"
    assert ranked[0]["id"] == "obs"


def test_exact_and_semantic_signals_merge_as_both():
    requirement = "OBS object store"
    products = _obs_catalog()
    embedder = _ScriptedEmbedder(
        [
            (("obs", "s3-api", "object store"), [1.0, 0.0, 0.0]),
            (("desk phone", "handset"), [0.0, 1.0, 0.0]),
        ]
    )
    product_vectors, requirement_vectors = _embed_catalog_and_requirements(
        embedder,
        products,
        [{"text": requirement}],
    )
    ranked = rank_requirement_products(
        requirement,
        products,
        product_vectors=product_vectors,
        requirement_vector=requirement_vectors[0],
    )
    obs = next(card for card in ranked if card["name"] == "OBS")
    assert obs["match_signal"] == "both"


class _DraftProvider:
    def __init__(self, text: str) -> None:
        self.text = text
        self.chunks = None

    def generate_grounded_answer(self, question, context_chunks, *, system_prompt, **kwargs):
        self.chunks = context_chunks
        from app.llm.base import GroundedAnswer

        return GroundedAnswer(
            text=self.text,
            citations=[],
            model_id="fake",
            prompt_version="test",
        )


def _azure_requirement() -> dict:
    return {
        "id": "R1",
        "text": "Provide Azure Blob Storage backups.",
        "evidence": [{"citation": "OBS datasheet p.1", "fenced_text": "OBS stores objects."}],
        "candidate_products": [
            {
                "id": "obs",
                "name": "OBS",
                "description": "S3-API-compatible object store with lifecycle policies",
                "match_signal": "semantic_match",
                "vendor": "Huawei",
            }
        ],
    }


def test_draft_one_uses_the_model_and_fences_requirement_and_evidence():
    provider = _DraftProvider(
        json.dumps(
            {
                "status": "Partial",
                "products": ["OBS"],
                "response": (
                    "OBS is functionally equivalent to Azure Blob Storage: "
                    "both are S3-API-compatible object storage with lifecycle policies."
                ),
                "citations": ["OBS datasheet p.1"],
                "confidence": 0.72,
            }
        )
    )
    drafted = _draft_one(_azure_requirement(), set(), provider=provider)
    assert drafted["status"] == "Partial"
    assert drafted["products"] == ["OBS"]
    assert drafted["citations"] == ["OBS datasheet p.1"]
    assert drafted["confidence"] == 0.72
    assert "functionally equivalent" in drafted["response"]
    assert drafted["used_semantic_match"] is True
    fenced = " ".join(chunk["fenced_text"] for chunk in provider.chunks)
    assert "UNTRUSTED_DOCUMENT_CONTENT" in fenced
    assert "Azure Blob Storage" in fenced
    assert "OBS datasheet p.1" in fenced


def test_draft_one_drops_invented_citations_and_falls_back():
    provider = _DraftProvider(
        json.dumps(
            {
                "status": "Compliant",
                "products": ["OBS"],
                "response": "Cited a document that was not retrieved.",
                "citations": ["Invented brief p.9"],
                "confidence": 0.99,
            }
        )
    )
    drafted = _draft_one(_azure_requirement(), set(), provider=provider)
    assert drafted["response"].startswith("Mapped to OBS")
    assert drafted["citations"] == ["OBS datasheet p.1"]
    assert drafted["used_semantic_match"] is True


def test_draft_one_falls_back_when_the_model_raises():
    class _Down:
        def generate_grounded_answer(self, *args, **kwargs):
            raise RuntimeError("model down")

    drafted = _draft_one(
        {"id": "R1", "text": "Secret unannounced discount?", "evidence": [], "candidate_products": []},
        set(),
        provider=_Down(),
    )
    assert drafted["status"] == "Non-Compliant"
    assert drafted["citations"] == []
    assert drafted["used_semantic_match"] is False


def test_draft_one_falls_back_when_the_token_budget_is_exhausted():
    class _Budget:
        def generate_grounded_answer(self, *args, **kwargs):
            raise TokenBudgetExhausted("monthly budget exhausted")

    drafted = _draft_one(
        {"id": "R1", "text": "Secret unannounced discount?", "evidence": [], "candidate_products": []},
        set(),
        provider=_Budget(),
    )
    assert drafted["status"] == "Non-Compliant"
    assert "sufficient information" in drafted["response"].lower()


def test_semantic_match_is_flagged_even_when_confidence_is_high():
    assert _needs_human_review({"confidence": 0.99, "used_semantic_match": True}) is True
    assert _needs_human_review({"confidence": 0.5, "used_semantic_match": False}) is True
    assert _needs_human_review({"confidence": 0.9, "unmet_prerequisites": True}) is True
    assert (
        _needs_human_review(
            {"confidence": 0.99, "unmet_prerequisites": False, "used_semantic_match": False}
        )
        is False
    )


def test_export_records_semantic_equivalence_for_review():
    data = render_docx(
        [
            {
                "text": "Provide Azure Blob Storage.",
                "response": "OBS is functionally equivalent to Azure Blob Storage.",
                "status": "Partial",
                "citations": ["OBS datasheet p.1"],
                "products": ["OBS"],
                "used_semantic_match": True,
            }
        ]
    )
    from docx import Document

    doc = Document(io.BytesIO(data))
    blob = " ".join(
        [paragraph.text for paragraph in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )
    assert "AI-inferred equivalence, human-reviewed" in blob
    assert "Provide Azure Blob Storage." in blob
    assert "OBS" in blob
