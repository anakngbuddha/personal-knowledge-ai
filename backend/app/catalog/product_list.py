"""3.3 Read a product list out of a spreadsheet.

Nobody types their catalogue into a form. They have it in a CSV or an XLSX that came
out of a distributor portal, with whatever column titles that portal felt like using.
So this module does three things, all of them pure:

* read the file into headers and rows (CSV or XLSX, first sheet, no pandas);
* guess which column means what, so the mapping screen opens already filled in;
* turn a mapping into product drafts, with a plain-language problem for every row it
  had to drop.

Nothing here touches the database, which is what makes the guessing and the
normalising testable without one. The writing half lives in ``importer.py``.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

_PUNCTUATION = re.compile(r"[^a-z0-9]+")
_ALIAS_SPLIT = re.compile(r"[,;|/]")

#: How many rows we are willing to skip looking for the column titles.
_MAX_HEADER_SCAN = 10


def normalize_header(value: str) -> str:
    return _PUNCTUATION.sub(" ", (value or "").strip().lower()).strip()


def normalize_name(value: str) -> str:
    """The same comparison key the map read uses, so both agree what one product is."""
    return _PUNCTUATION.sub(" ", (value or "").strip().lower()).strip()


@dataclass(frozen=True)
class ImportField:
    """One thing a column can fill in."""

    key: str
    label: str
    aliases: tuple[str, ...] = ()
    required: bool = False
    note: str = ""


FIELDS: tuple[ImportField, ...] = (
    ImportField(
        "name",
        "Product name",
        ("product", "product name", "model", "model name", "sku", "item", "part", "part number"),
        required=True,
        note="The one column we cannot do without.",
    ),
    ImportField("vendor", "Brand", ("brand", "manufacturer", "make", "supplier", "oem")),
    ImportField("category", "What it is", ("type", "product type", "family", "group", "line")),
    ImportField(
        "ownership",
        "Ours or resold",
        ("own or resold", "source", "resold", "owned", "own"),
    ),
    ImportField("tier", "Tier", ("level", "range")),
    ImportField("deployment_model", "Where it runs", ("deployment", "hosting", "cloud or on prem")),
    ImportField("licensing_model", "How it is licensed", ("licensing", "licence", "license")),
    ImportField("target_segment", "Who it is for", ("segment", "market", "audience")),
    ImportField("lifecycle_status", "Still sold", ("lifecycle", "status", "availability", "eol")),
    ImportField("description", "Description", ("summary", "notes", "details", "about")),
    ImportField("prerequisites", "What it needs", ("requires", "needs", "prerequisite")),
    ImportField("support_path", "Support", ("support path", "warranty")),
    ImportField("aliases", "Also called", ("alias", "also known as", "aka", "other names")),
    ImportField("partner_tier", "Partner tier", ("partner level", "reseller tier")),
    ImportField("margin_band", "Margin", ("margin band", "discount")),
    ImportField("support_owner", "Who supports it", ("support by", "supported by")),
    ImportField("source_of_truth_url", "Link", ("url", "link", "datasheet", "website")),
)

FIELDS_BY_KEY = {item.key: item for item in FIELDS}

#: Everything we are prepared to accept for the fields the database constrains. An
#: unrecognised value falls back to the safe default rather than failing the row: a
#: spreadsheet full of local spellings should still import.
_OWNERSHIP = {
    "own": "own",
    "ours": "own",
    "owned": "own",
    "internal": "own",
    "in house": "own",
    "our product": "own",
    "first party": "own",
    "resold": "resold",
    "resell": "resold",
    "reseller": "resold",
    "partner": "resold",
    "third party": "resold",
    "distributed": "resold",
    "vendor": "resold",
}

_DEPLOYMENT = {
    "cloud": "cloud",
    "saas": "cloud",
    "hosted": "cloud",
    "on prem": "on-prem",
    "on premise": "on-prem",
    "on premises": "on-prem",
    "onprem": "on-prem",
    "local": "on-prem",
    "appliance": "on-prem",
    "hardware": "on-prem",
    "hybrid": "hybrid",
    "both": "hybrid",
}

_LIFECYCLE = {
    "ga": "GA",
    "generally available": "GA",
    "available": "GA",
    "current": "GA",
    "shipping": "GA",
    "active": "GA",
    "yes": "GA",
    "eol": "EOL",
    "end of life": "EOL",
    "end of sale": "EOL",
    "discontinued": "EOL",
    "retired": "EOL",
    "no": "EOL",
    "roadmap": "roadmap",
    "future": "roadmap",
    "announced": "roadmap",
    "coming": "roadmap",
    "planned": "roadmap",
}

_LENGTH_LIMITS = {
    "name": 512,
    "vendor": 255,
    "category": 128,
    "tier": 64,
    "licensing_model": 64,
    "target_segment": 64,
    "partner_tier": 128,
    "margin_band": 64,
    "support_owner": 64,
}


@dataclass
class Table:
    headers: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    sheet_name: str | None = None
    truncated: bool = False

    @property
    def row_count(self) -> int:
        return len(self.rows)


@dataclass
class ProductDraft:
    row_number: int
    name: str
    values: dict = field(default_factory=dict)

    @property
    def key(self) -> str:
        return normalize_name(self.name)


@dataclass
class RowProblem:
    row_number: int
    reason: str

    def as_dict(self) -> dict:
        return {"row": self.row_number, "reason": self.reason}


@dataclass
class ImportPlan:
    drafts: list[ProductDraft] = field(default_factory=list)
    problems: list[RowProblem] = field(default_factory=list)
    mapping: dict = field(default_factory=dict)
    truncated: bool = False

    @property
    def count(self) -> int:
        return len(self.drafts)


class UnsupportedProductList(Exception):
    """The file is not something we can read as a list. Carries wording for a person."""


# -- reading ---------------------------------------------------------------


def read_table(data: bytes, *, filename: str, max_rows: int = 2000) -> Table:
    suffix = (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    if suffix in {"csv", "tsv", "txt"}:
        return _read_delimited(data, max_rows=max_rows)
    if suffix in {"xlsx", "xlsm"}:
        return _read_xlsx(data, max_rows=max_rows)
    raise UnsupportedProductList(
        "Save the list as a CSV or an Excel file (.xlsx) and try again."
    )


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _read_delimited(data: bytes, *, max_rows: int) -> Table:
    text = _decode(data)
    sample = text[:4096]
    delimiter = ","
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        lines = sample.splitlines()
        if lines and "\t" in lines[0]:
            delimiter = "\t"
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return _from_rows([[_cell(cell) for cell in row] for row in reader], max_rows=max_rows)


def _read_xlsx(data: bytes, *, max_rows: int) -> Table:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - openpyxl ships with the backend
        raise UnsupportedProductList("Save the list as a CSV and try again.") from exc

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - a broken file is a user problem, not a crash
        raise UnsupportedProductList(
            "That spreadsheet could not be opened. Try saving it again as .xlsx or CSV."
        ) from exc
    try:
        sheet = workbook.worksheets[0]
        raw: list[list[str]] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            raw.append([_cell(value) for value in row])
            if index > max_rows + _MAX_HEADER_SCAN:
                break
        table = _from_rows(raw, max_rows=max_rows)
        table.sheet_name = sheet.title
        return table
    finally:
        workbook.close()


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _from_rows(raw: list[list[str]], *, max_rows: int) -> Table:
    header_index = None
    for index, row in enumerate(raw[:_MAX_HEADER_SCAN]):
        if any(cell for cell in row):
            header_index = index
            break
    if header_index is None:
        raise UnsupportedProductList("That file has no rows in it.")

    headers = list(raw[header_index])
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise UnsupportedProductList("That file has no column titles in its first row.")

    body: list[list[str]] = []
    truncated = False
    for row in raw[header_index + 1 :]:
        if not any(cell for cell in row):
            continue
        if len(body) >= max_rows:
            truncated = True
            break
        body.append(list(row[: len(headers)]) + [""] * max(0, len(headers) - len(row)))

    return Table(headers=headers, rows=body, truncated=truncated)


# -- guessing --------------------------------------------------------------


def suggest_mapping(headers: list[str]) -> dict:
    """Guess a field for each column. An empty string means "leave this one out".

    Keys are column positions as strings, not titles: a spreadsheet is allowed to use
    the same title twice, and position is the only thing that is always unique.
    """
    lookup: dict[str, str] = {}
    for item in FIELDS:
        lookup[normalize_header(item.label)] = item.key
        lookup[normalize_header(item.key)] = item.key
        for alias in item.aliases:
            lookup.setdefault(normalize_header(alias), item.key)

    mapping: dict[str, str] = {}
    used: set[str] = set()
    for index, header in enumerate(headers):
        cleaned = normalize_header(header)
        guess = lookup.get(cleaned, "")
        if not guess:
            for candidate, key in lookup.items():
                if candidate and len(candidate) > 3 and candidate in cleaned:
                    guess = key
                    break
        # One column per field. A second candidate is left for the person to decide.
        if guess and guess in used:
            guess = ""
        if guess:
            used.add(guess)
        mapping[str(index)] = guess
    return mapping


# -- planning --------------------------------------------------------------


def build_plan(table: Table, mapping: dict) -> ImportPlan:
    """Turn a mapping into product drafts, dropping rows nobody could use."""
    columns: list[tuple[int, str]] = []
    for raw_index, key in (mapping or {}).items():
        if not key or key not in FIELDS_BY_KEY:
            continue
        try:
            index = int(raw_index)
        except (TypeError, ValueError):
            continue
        if 0 <= index < len(table.headers):
            columns.append((index, key))

    plan = ImportPlan(mapping=dict(mapping or {}), truncated=table.truncated)
    if not any(key == "name" for _, key in columns):
        plan.problems.append(
            RowProblem(0, "Pick the column that holds the product name, then import again.")
        )
        return plan

    seen: dict[str, int] = {}
    for offset, row in enumerate(table.rows):
        row_number = offset + 1
        values: dict = {}
        for index, key in columns:
            cell = row[index] if index < len(row) else ""
            cleaned = _clean_value(key, cell)
            if cleaned not in (None, "", []):
                values[key] = cleaned

        name = str(values.pop("name", "")).strip()
        if not name:
            plan.problems.append(RowProblem(row_number, "No product name in this row."))
            continue
        key = normalize_name(name)
        if key in seen:
            plan.problems.append(
                RowProblem(row_number, f"Same product as row {seen[key]}, so it was skipped.")
            )
            continue
        seen[key] = row_number
        plan.drafts.append(ProductDraft(row_number=row_number, name=name[:512], values=values))

    return plan


def _clean_value(key: str, raw: str):
    value = (raw or "").strip()
    if not value:
        return ""
    if key == "ownership":
        return _OWNERSHIP.get(normalize_header(value), "own")
    if key == "deployment_model":
        return _DEPLOYMENT.get(normalize_header(value), "cloud")
    if key == "lifecycle_status":
        return _LIFECYCLE.get(normalize_header(value), "GA")
    if key == "aliases":
        parts = [part.strip() for part in _ALIAS_SPLIT.split(value)]
        return [part[:255] for part in parts if part][:6]
    if key in _LENGTH_LIMITS:
        return value[: _LENGTH_LIMITS[key]]
    return value[:4000]


def field_catalogue() -> list[dict]:
    """What the mapping screen offers, in plain words."""
    return [
        {"key": item.key, "label": item.label, "required": item.required, "note": item.note}
        for item in FIELDS
    ]
