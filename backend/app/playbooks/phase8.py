"""Phase 8 playbooks: solution composition, incident triage, and upgrade impact."""

from __future__ import annotations

import io
import re
from typing import Any

from app.catalog.service import CatalogService
from app.documents.injection import neutralize_fences, scan_for_injection
from app.storage.factory import get_storage
from app.tools.catalog_tools import catalog_impact
from app.tools.registry import ToolContext
from app.tools.search_tools import hybrid_search_tool
from app.workflows.handlers import HandlerContext, register


def _tokens(value: str) -> set[str]:
    return {token for token in re.findall(r"[a-z0-9][a-z0-9+.-]{2,}", value.lower())}


def extract_discovery_constraints(notes: str, vendors: list[str] | None = None) -> dict[str, Any]:
    """Extract deterministic, auditable constraints without treating notes as instructions."""
    clean = neutralize_fences(notes)[:50_000]
    lowered = clean.lower()
    deployment = None
    if re.search(r"\bon[- ]prem(?:ises)?(?: only)?\b", lowered):
        deployment = "on-prem"
    elif re.search(r"\bhybrid(?: only)?\b", lowered):
        deployment = "hybrid"
    elif re.search(r"\bcloud(?: only| native)?\b", lowered):
        deployment = "cloud"
    budget_match = re.search(r"(?:budget|cap(?:ped)? at)\s*[:=]?\s*([$€£]?\s*[\d,.]+(?:\s*[kKmM])?)", clean, re.I)
    user_match = re.search(r"([\d,]+)\s+(?:users|seats|employees|endpoints)\b", clean, re.I)
    forbidden: list[str] = []
    for vendor in vendors or []:
        pattern = rf"\b(?:avoid|exclude|no|not)\s+{re.escape(vendor.lower())}\b"
        if re.search(pattern, lowered):
            forbidden.append(vendor)
    return {
        "deployment_model": deployment,
        "budget": budget_match.group(1).strip() if budget_match else None,
        "scale": int(user_match.group(1).replace(",", "")) if user_match else None,
        "forbidden_vendors": forbidden,
        "keywords": sorted(_tokens(clean))[:100],
        "injection_findings": [finding.as_dict() for finding in scan_for_injection(clean)],
    }


def _product_payload(product: Any, impact: dict[str, Any], score: int) -> dict[str, Any]:
    return {
        "id": str(product.id),
        "name": product.name,
        "slug": product.slug,
        "vendor": product.vendor,
        "category": product.category,
        "deployment_model": product.deployment_model,
        "licensing_model": product.licensing_model,
        "description": product.description,
        "score": score,
        "impact": impact,
    }


def select_compatible_bundle(candidates: list[dict[str, Any]], limit: int = 8) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Greedily select a scored bundle while refusing graph conflicts."""
    chosen: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    selected_names: set[str] = set()
    for candidate in sorted(candidates, key=lambda row: (-int(row.get("score", 0)), str(row.get("name", "")))):
        incompatibilities = {
            str(item.get("conflicted_product_name"))
            for item in (candidate.get("impact") or {}).get("all_incompatibilities") or []
            if item.get("conflicted_product_name")
        }
        conflict = sorted(selected_names & incompatibilities)
        if conflict:
            rejected.append({"product": candidate.get("name"), "reason": "conflicts_with", "conflicts": conflict})
            continue
        chosen.append(candidate)
        selected_names.add(str(candidate.get("name")))
        if len(chosen) >= limit:
            break
    return chosen, rejected


def _docx(title: str, sections: list[tuple[str, str]], table: tuple[list[str], list[list[str]]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    for heading, body in sections:
        document.add_heading(heading, level=2)
        document.add_paragraph(body or "Not available from approved evidence.")
    if table:
        headers, rows = table
        document.add_heading("Bill of Materials", level=2)
        grid = document.add_table(rows=1, cols=len(headers))
        grid.style = "Table Grid"
        for idx, header in enumerate(headers):
            grid.rows[0].cells[idx].text = header
        for values in rows:
            cells = grid.add_row().cells
            for idx, value in enumerate(values):
                cells[idx].text = value
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _store_docx(ctx: HandlerContext, filename: str, data: bytes) -> dict[str, Any]:
    key = f"phase8/{ctx.run.org_id}/{ctx.run.id}/{filename}"
    get_storage().put(
        key,
        data,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    return {"storage_key": key, "filename": filename, "bytes": len(data)}


@register("solution.extract_constraints")
def solution_extract_constraints(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    source = ctx.run.input_payload or {}
    notes = str(source.get("notes") or "")
    products = CatalogService(ctx.db).list_products(ctx.run.workspace_id, limit=500)
    return {
        "notes_excerpt": neutralize_fences(notes)[:2000],
        "constraints": extract_discovery_constraints(notes, sorted({p.vendor for p in products})),
    }


@register("solution.compose_candidates")
def solution_compose_candidates(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    extracted = ctx.upstream.get("extract_constraints") or {}
    constraints = dict(extracted.get("constraints") or {})
    keywords = set(constraints.get("keywords") or [])
    service = CatalogService(ctx.db)
    products = service.list_products(ctx.run.workspace_id, limit=500)
    tool_ctx = ToolContext(ctx.db, ctx.principal, ctx.run.workspace_id)
    forbidden = {str(v).lower() for v in constraints.get("forbidden_vendors") or []}
    requested_deployment = constraints.get("deployment_model")
    candidates: list[dict[str, Any]] = []
    for product in products:
        if product.lifecycle_status != "GA" or product.vendor.lower() in forbidden:
            continue
        if requested_deployment and product.deployment_model not in {requested_deployment, "hybrid"}:
            continue
        haystack = _tokens(
            f"{product.name} {product.vendor} {product.category} {product.description or ''} "
            + " ".join(pc.capability.name for pc in product.capabilities if pc.capability)
        )
        score = len(keywords & haystack)
        if keywords and score == 0:
            continue
        impact = catalog_impact(tool_ctx, product.name)
        candidates.append(_product_payload(product, impact if "error" not in impact else {}, score))
    selected, rejected = select_compatible_bundle(candidates)
    return {"constraints": constraints, "candidates": candidates[:30], "selected": selected, "rejected": rejected}


@register("solution.validate_bundle")
def solution_validate_bundle(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    composed = ctx.upstream.get("compose_candidates") or {}
    selected = list(composed.get("selected") or [])
    selected_names = {str(row.get("name")) for row in selected}
    warnings: list[dict[str, Any]] = []
    prerequisites: list[dict[str, Any]] = []
    for row in selected:
        impact = row.get("impact") or {}
        for item in impact.get("all_prerequisites") or []:
            prerequisites.append({"product": row.get("name"), **item, "included": item.get("name") in selected_names})
        for item in impact.get("all_incompatibilities") or []:
            if item.get("conflicted_product_name") in selected_names:
                warnings.append({"product": row.get("name"), **item})
    return {
        **composed,
        "prerequisites": prerequisites,
        "warnings": warnings,
        "valid": bool(selected) and not warnings,
    }


@register("solution.generate_hld")
def solution_generate_hld(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    validated = ctx.upstream.get("validate_bundle") or {}
    selected = list(validated.get("selected") or [])
    constraints = dict(validated.get("constraints") or {})
    bom = [
        {
            "product": row.get("name"),
            "vendor": row.get("vendor"),
            "category": row.get("category"),
            "deployment": row.get("deployment_model"),
            "quantity": constraints.get("scale") or 1,
            "licensing": row.get("licensing_model"),
        }
        for row in selected
    ]
    product_names = ", ".join(str(row.get("name")) for row in selected) or "No evidence-backed bundle available"
    hld = {
        "summary": f"Proposed graph-validated solution: {product_names}.",
        "deployment": constraints.get("deployment_model") or "customer choice",
        "architecture": [
            {"component": row.get("name"), "role": row.get("category"), "vendor": row.get("vendor")}
            for row in selected
        ],
        "prerequisites": validated.get("prerequisites") or [],
        "warnings": validated.get("warnings") or [],
    }
    review = {
        "id": "HLD-1",
        "text": "Review the proposed HLD and Bill of Materials",
        "response": hld["summary"],
        "status": "Compliant" if validated.get("valid") else "Partially",
        "confidence": 0.9 if validated.get("valid") else 0.5,
        "products": [row.get("name") for row in selected],
        "citations": [
            item.get("evidence") for row in selected for item in (row.get("impact") or {}).get("all_prerequisites") or [] if item.get("evidence")
        ],
        "unmet_prerequisites": any(not item.get("included") for item in validated.get("prerequisites") or []),
    }
    return {"hld": hld, "bom": bom, "answers": [review], "needs_review": True}


@register("solution.gate")
def solution_gate(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    return dict(ctx.upstream.get("generate_hld") or {})


@register("solution.export")
def solution_export(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    approved = ctx.upstream.get("human_gate") or {}
    hld = approved.get("hld") or {}
    bom = list(approved.get("bom") or [])
    data = _docx(
        "High-Level Design",
        [
            ("Executive Summary", str(hld.get("summary") or "")),
            ("Deployment Model", str(hld.get("deployment") or "")),
            ("Architecture", "\n".join(f"{x.get('component')}: {x.get('role')}" for x in hld.get("architecture") or [])),
            ("Prerequisites", "\n".join(str(x.get("name")) for x in hld.get("prerequisites") or [])),
        ],
        (["Product", "Vendor", "Category", "Deployment", "Quantity", "Licensing"], [
            [str(row.get(k) or "") for k in ("product", "vendor", "category", "deployment", "quantity", "licensing")]
            for row in bom
        ]),
    )
    return {**_store_docx(ctx, "solution-hld.docx", data), "product_count": len(bom)}


@register("incident.analyze")
def incident_analyze(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    source = ctx.run.input_payload or {}
    logs = neutralize_fences(str(source.get("logs") or ""))[:50_000]
    lines = [line.strip() for line in logs.splitlines() if line.strip()]
    signals = [line for line in lines if re.search(r"error|fail|timeout|exception|critical|denied|unavailable", line, re.I)]
    return {
        "signals": signals[:100],
        "keywords": sorted(_tokens(" ".join(signals or lines)))[:100],
        "injection_findings": [finding.as_dict() for finding in scan_for_injection(logs)],
    }


@register("incident.traverse")
def incident_traverse(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    source = ctx.run.input_payload or {}
    requested = {str(name).lower() for name in source.get("install_base") or []}
    service = CatalogService(ctx.db)
    products = service.list_products(ctx.run.workspace_id, limit=500)
    tool_ctx = ToolContext(ctx.db, ctx.principal, ctx.run.workspace_id)
    matched = [p for p in products if p.name.lower() in requested or p.slug.lower() in requested]
    impacts = [catalog_impact(tool_ctx, product.name) for product in matched]
    return {"products": [p.name for p in matched], "impacts": impacts, "signals": (ctx.upstream.get("analyze_logs") or {}).get("signals") or []}


@register("incident.retrieve")
def incident_retrieve(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    analyzed = ctx.upstream.get("analyze_logs") or {}
    query = " ".join(analyzed.get("keywords") or [])[:2000]
    hits: list[dict[str, Any]] = []
    if query:
        try:
            hits = (hybrid_search_tool(ToolContext(ctx.db, ctx.principal, ctx.run.workspace_id), query).get("hits") or [])[:8]
        except Exception:  # sqlite and deployments without pgvector still produce a graph-only runbook
            hits = []
    return {"evidence": hits}


@register("incident.build_runbook")
def incident_build_runbook(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    traversed = ctx.upstream.get("traverse_install_base") or {}
    evidence = (ctx.upstream.get("retrieve_evidence") or {}).get("evidence") or []
    steps = [
        "Capture timestamps, affected users, and the first reproducible failure.",
        "Validate health and connectivity for every installed dependency listed below.",
        "Check direct and transitive prerequisites before restarting or upgrading components.",
        "Compare the observed signal with the cited approved collateral.",
        "Escalate with the evidence bundle if the fault remains unresolved.",
    ]
    return {
        "title": "Incident Root-Cause Triage Runbook",
        "products": traversed.get("products") or [],
        "signals": traversed.get("signals") or [],
        "impacts": traversed.get("impacts") or [],
        "steps": steps,
        "citations": [hit.get("citation") for hit in evidence if hit.get("citation")],
    }


@register("incident.export")
def incident_export(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    runbook = ctx.upstream.get("build_runbook") or {}
    data = _docx(
        str(runbook.get("title") or "Incident Runbook"),
        [
            ("Affected Products", ", ".join(runbook.get("products") or [])),
            ("Observed Signals", "\n".join(runbook.get("signals") or [])),
            ("Troubleshooting Steps", "\n".join(f"{i + 1}. {step}" for i, step in enumerate(runbook.get("steps") or []))),
            ("Citations", "\n".join(runbook.get("citations") or [])),
        ],
    )
    return {**_store_docx(ctx, "incident-triage-runbook.docx", data), "step_count": len(runbook.get("steps") or [])}


@register("upgrade.resolve")
def upgrade_resolve(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    product = str((ctx.run.input_payload or {}).get("product") or "")
    impact = catalog_impact(ToolContext(ctx.db, ctx.principal, ctx.run.workspace_id), product)
    return {"requested_product": product, "impact": impact}


@register("upgrade.audit")
def upgrade_audit(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    resolved = ctx.upstream.get("resolve_product") or {}
    impact = resolved.get("impact") or {}
    return {
        "product": impact.get("product_name") or resolved.get("requested_product"),
        "proposed_version": (ctx.run.input_payload or {}).get("proposed_version"),
        "prerequisites": impact.get("all_prerequisites") or [],
        "breaking_impacts": impact.get("all_incompatibilities") or [],
        "integrations": impact.get("direct_integrations") or [],
        "alternatives": impact.get("alternatives") or [],
        "unknown_product": impact.get("error") == "unknown_product",
    }


@register("upgrade.export")
def upgrade_export(ctx: HandlerContext, payload: dict[str, Any]) -> dict[str, Any]:
    audit = ctx.upstream.get("audit_impact") or {}
    data = _docx(
        "Upgrade Impact Audit",
        [
            ("Product", f"{audit.get('product')} {audit.get('proposed_version') or ''}".strip()),
            ("Prerequisites", "\n".join(str(x.get("name")) for x in audit.get("prerequisites") or [])),
            ("Breaking Impacts", "\n".join(f"{x.get('conflicted_product_name')}: {x.get('reason')}" for x in audit.get("breaking_impacts") or [])),
            ("Integrations", "\n".join(str(x.get("name")) for x in audit.get("integrations") or [])),
            ("Alternatives", "\n".join(str(x.get("name")) for x in audit.get("alternatives") or [])),
        ],
    )
    return {**_store_docx(ctx, "upgrade-impact-audit.docx", data), "breaking_impact_count": len(audit.get("breaking_impacts") or [])}
