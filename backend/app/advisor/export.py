"""4.2 Export a recommendation to Word or PowerPoint.

Both writers use libraries the backend already depends on (``python-docx``,
``python-pptx``), so nothing new has to be installed and nothing native has to be
present on Render.

The two formats exist because they are used at different moments: the Word file is
what gets attached to an email after a call, the deck is what gets shown during one.
Both are built from the same :class:`Recommendation`, so they cannot drift.

Provenance survives the export. A product with no document behind it is labelled
"General knowledge" in the table and repeated under "Verify before you send this",
because a recommendation that loses its caveats on the way into a .docx is worse than
no export at all.
"""

from __future__ import annotations

import io

from app.advisor.recommend import Recommendation

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

#: A slide that needs a scrollbar is not a slide.
_BULLETS_PER_SLIDE = 6


def filename_for(recommendation: Recommendation, extension: str) -> str:
    who = (recommendation.customer or "customer").lower()
    safe = "".join(char if char.isalnum() or char in "-_" else "-" for char in who).strip("-")
    return f"recommendation-{safe or 'customer'}.{extension}"


def to_docx(recommendation: Recommendation) -> bytes:
    """A proposal-shaped Word document: prose, then the table, then the caveats."""
    from docx import Document

    document = Document()
    document.add_heading(
        f"Recommendation for {recommendation.customer or 'the customer'}", level=1
    )

    document.add_heading("What they asked for", level=2)
    if recommendation.restated:
        for item in recommendation.restated:
            document.add_paragraph(item, style="List Bullet")
    else:
        document.add_paragraph("No specific requirements were captured.")

    document.add_heading("What we would offer", level=2)
    if not recommendation.lines:
        document.add_paragraph(
            "Nothing in the current product list matches these requirements. "
            "Add the products you sell and run this again."
        )
    for line in recommendation.lines:
        document.add_heading(line.requirement, level=3)
        for pick in line.picks:
            document.add_paragraph(f"{pick.name} ({pick.vendor}) - {pick.why()}", style="List Bullet")

    headers, rows = recommendation.table()
    if rows:
        document.add_heading("At a glance", level=2)
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = str(value)

    for title, items in _closing_sections(recommendation):
        if not items:
            continue
        document.add_heading(title, level=2)
        for item in items:
            document.add_paragraph(item, style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_pptx(recommendation: Recommendation) -> bytes:
    """A short deck: a title, the requirements, a slide per requirement, the caveats."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    title_layout = presentation.slide_layouts[0]
    bullet_layout = presentation.slide_layouts[1]

    opening = presentation.slides.add_slide(title_layout)
    opening.shapes.title.text = f"Recommendation for {recommendation.customer or 'the customer'}"
    opening.placeholders[1].text = (
        f"{len(recommendation.picks)} products across {len(recommendation.lines)} requirements"
    )

    _bullet_slide(
        presentation,
        bullet_layout,
        "What they asked for",
        recommendation.restated or ["No specific requirements were captured."],
    )

    for line in recommendation.lines:
        _bullet_slide(
            presentation,
            bullet_layout,
            line.requirement[:120],
            [f"{pick.name} ({pick.vendor}) - {pick.why()}" for pick in line.picks],
        )

    if recommendation.picks:
        _table_slide(presentation, recommendation, Inches, Pt)

    for title, items in _closing_sections(recommendation):
        if items:
            _bullet_slide(presentation, bullet_layout, title, items)

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


# -- shared ----------------------------------------------------------------


def _closing_sections(recommendation: Recommendation) -> list[tuple[str, list[str]]]:
    return [
        (
            "Fits and clashes",
            [
                f"{row['product']} clashes with {row['clashes_with']}: {row['reason']}"
                for row in recommendation.conflicts
            ],
        ),
        (
            "Also needed to make this work",
            [
                f"{row['needed']} (for {row['product']})"
                + ("" if row["in_product_list"] else " - not in your product list")
                for row in recommendation.also_needed
            ],
        ),
        ("Not covered by your product list", recommendation.gaps),
        (
            "Worth adding to the deal",
            [
                f"{row['product']} - {row['kind']} of {row['from_product']}"
                for row in recommendation.growth
            ],
        ),
        ("Ask the customer", recommendation.questions),
        ("Verify before you send this", recommendation.assumptions),
    ]


def _bullet_slide(presentation, layout, title: str, bullets: list[str]) -> None:
    """One slide, or several when the list is long. Never an overflowing one."""
    chunks = [
        bullets[index : index + _BULLETS_PER_SLIDE]
        for index in range(0, max(len(bullets), 1), _BULLETS_PER_SLIDE)
    ] or [[]]
    for position, chunk in enumerate(chunks):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title if position == 0 else f"{title} (cont.)"
        frame = slide.placeholders[1].text_frame
        frame.clear()
        for index, bullet in enumerate(chunk):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = str(bullet)[:300]


def _table_slide(presentation, recommendation: Recommendation, Inches, Pt) -> None:
    headers, rows = recommendation.table()
    rows = rows[:10]  # a slide is a summary; the Word file carries the whole list
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "At a glance"
    shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(0.4), Inches(1.6), Inches(9.2), Inches(0.8)
    )
    table = shape.table
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    for row_index, row in enumerate(rows, start=1):
        for column, value in enumerate(row):
            cell = table.cell(row_index, column)
            cell.text = str(value)[:80]
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
